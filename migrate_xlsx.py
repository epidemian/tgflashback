"""One-off import of Chanchullo Flashback.xlsx into SQLite.

Idempotent: reruns just upsert the same (player_code, puzzle_date) rows.

Date anchoring:
- `Puntajes` (2026): the `Missing Se` sheet lists a date for each week index
  in the same row order (row 2 of both sheets = week index 1), so we zip the
  two sheets by row position. Those recorded dates are all Fridays, but the
  group actually plays/shares on Saturday (confirmed against a real
  NYT-parsed "Flashback for ..." message) — so we add one day to land on the
  real play date, e.g. row 2's recorded 2026-01-02 (Fri) -> 2026-01-03 (Sat).
- `Puntajes2025`: has no matching "Missing" sheet with real dates. We assume
  its last row (Fecha=18) is the week immediately before `Puntajes` row 1
  (2026-01-03 - 7 days = 2025-12-27) and step back 7 days per row from there.
  This is an approximation accepted in absence of an exact recorded date.
"""

import datetime

import openpyxl

from bot import db
from bot.config import PLAYER_CODES

XLSX_PATH = "Chanchullo Flashback.xlsx"

DAY_CORRECTION = datetime.timedelta(days=1)  # sheet records Friday; real play day is Saturday

PUNTAJES_2026_ANCHOR = datetime.date(2026, 1, 2) + DAY_CORRECTION  # Missing Se row 2 (Fecha=1)
PUNTAJES_2025_LAST_ROW_DATE = PUNTAJES_2026_ANCHOR - datetime.timedelta(days=7)


def _score_columns(ws) -> dict[int, str]:
    """Map column index (1-based) -> player code, from the header row."""
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    columns = {}
    for idx, value in enumerate(header, start=1):
        if value in PLAYER_CODES:
            columns[idx] = value
    return columns


def import_puntajes_2026(wb) -> int:
    puntajes = wb["Puntajes"]
    missing = wb["Missing Se"]
    columns = _score_columns(puntajes)

    puntajes_rows = list(puntajes.iter_rows(min_row=2, values_only=True))
    missing_rows = list(missing.iter_rows(min_row=2, values_only=True))

    count = 0
    for p_row, m_row in zip(puntajes_rows, missing_rows):
        if p_row[0] is None:
            continue
        puzzle_date = m_row[0]
        if not isinstance(puzzle_date, datetime.datetime):
            continue
        puzzle_date_iso = (puzzle_date.date() + DAY_CORRECTION).isoformat()

        for col_idx, code in columns.items():
            score = p_row[col_idx - 1]
            if score is None:
                continue
            db.upsert_score(player_code=code, puzzle_date=puzzle_date_iso, score=int(score))
            count += 1
    return count


def import_puntajes_2025(wb) -> int:
    ws = wb["Puntajes2025"]
    columns = _score_columns(ws)
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    last_fecha = max(int(r[0]) for r in rows if r[0] is not None)

    count = 0
    for row in rows:
        if row[0] is None:
            continue
        fecha_idx = int(row[0])
        puzzle_date = PUNTAJES_2025_LAST_ROW_DATE - datetime.timedelta(
            days=7 * (last_fecha - fecha_idx)
        )
        puzzle_date_iso = puzzle_date.isoformat()

        for col_idx, code in columns.items():
            score = row[col_idx - 1]
            if score is None:
                continue
            db.upsert_score(player_code=code, puzzle_date=puzzle_date_iso, score=int(score))
            count += 1
    return count


def main() -> None:
    db.init_db()
    db.seed_players()

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)

    n2025 = import_puntajes_2025(wb)
    n2026 = import_puntajes_2026(wb)

    print(f"Importados {n2025} puntajes de 2025 y {n2026} de 2026.")


if __name__ == "__main__":
    main()
